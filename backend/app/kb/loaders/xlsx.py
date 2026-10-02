"""
Q2 Knowledge Base — XLSX / CSV Extractor.
"""

from __future__ import annotations

import csv
from pathlib import Path

from app.core.logging import get_logger
from app.kb.loaders.base import BaseExtractor
from app.kb.models import (
    ExtractedDocument,
    ExtractedElement,
    ExtractedTable,
    ExtractionDiagnostics,
    ExtractionStatus,
    ElementType,
    SourceRecord,
)

logger = get_logger(__name__)


class XLSXExtractor(BaseExtractor):
    """Extracts tables from XLSX files using openpyxl."""

    def supported_types(self) -> list[str]:
        return [".xlsx"]

    def extract(self, source: SourceRecord, filepath: Path) -> ExtractedDocument:
        doc = ExtractedDocument(source_id=source.source_id)
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(filepath), read_only=True, data_only=True)
            order = 0

            for sheet in wb.sheetnames:
                ws = wb[sheet]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                columns = [str(c) if c else f"col_{i}" for i, c in enumerate(rows[0])]
                data_rows = [[str(cell) if cell is not None else "" for cell in row] for row in rows[1:]]

                ext_table = ExtractedTable(
                    source_id=source.source_id,
                    columns=columns,
                    rows=data_rows,
                    section=sheet,
                )
                doc.tables.append(ext_table)

                table_elem = ExtractedElement(
                    source_id=source.source_id,
                    element_type=ElementType.TABLE,
                    text=ext_table.to_text(),
                    order=order,
                    section=sheet,
                    table_id=ext_table.table_id,
                )
                doc.elements.append(table_elem)
                order += 1

            wb.close()
            doc.raw_text = "\n\n".join(e.text for e in doc.elements)
            doc.diagnostics = ExtractionDiagnostics(
                total_characters=len(doc.raw_text),
                total_elements=len(doc.elements),
                tables_detected=len(doc.tables),
                status=ExtractionStatus.SUCCESS if doc.elements else ExtractionStatus.FAILED,
            )

        except Exception as e:
            logger.error("XLSX extraction failed", extra={"extra_data": {"source_id": source.source_id, "error": str(e)}})
            doc.diagnostics = ExtractionDiagnostics(status=ExtractionStatus.FAILED, issues=[f"XLSX error: {str(e)}"])

        return doc


class CSVExtractor(BaseExtractor):
    """Extracts tables from CSV files."""

    def supported_types(self) -> list[str]:
        return [".csv"]

    def extract(self, source: SourceRecord, filepath: Path) -> ExtractedDocument:
        doc = ExtractedDocument(source_id=source.source_id)
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = list(reader)

            if not rows:
                doc.diagnostics = ExtractionDiagnostics(status=ExtractionStatus.FAILED, issues=["Empty CSV file"])
                return doc

            columns = rows[0]
            data_rows = rows[1:]

            ext_table = ExtractedTable(
                source_id=source.source_id,
                columns=columns,
                rows=data_rows,
            )
            doc.tables.append(ext_table)

            table_elem = ExtractedElement(
                source_id=source.source_id,
                element_type=ElementType.TABLE,
                text=ext_table.to_text(),
                order=0,
                table_id=ext_table.table_id,
            )
            doc.elements.append(table_elem)

            doc.raw_text = table_elem.text
            doc.diagnostics = ExtractionDiagnostics(
                total_characters=len(doc.raw_text),
                total_elements=1,
                tables_detected=1,
                status=ExtractionStatus.SUCCESS,
            )

        except Exception as e:
            logger.error("CSV extraction failed", extra={"extra_data": {"source_id": source.source_id, "error": str(e)}})
            doc.diagnostics = ExtractionDiagnostics(status=ExtractionStatus.FAILED, issues=[f"CSV error: {str(e)}"])

        return doc
