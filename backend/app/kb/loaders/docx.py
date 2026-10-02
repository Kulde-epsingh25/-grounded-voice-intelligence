"""
Q2 Knowledge Base — DOCX Extractor.
"""

from __future__ import annotations

from pathlib import Path

from app.core.logging import get_logger
from app.kb.loaders.base import BaseExtractor
from app.kb.models import (
    ExtractedDocument,
    ExtractedElement,
    ExtractedTable,
    ExtractionDiagnostics,
    ExtractionStatus,
    ElementType,
    SourceRecord,
)

logger = get_logger(__name__)


class DOCXExtractor(BaseExtractor):
    """Extracts text and tables from DOCX files."""

    def supported_types(self) -> list[str]:
        return [".docx"]

    def extract(self, source: SourceRecord, filepath: Path) -> ExtractedDocument:
        doc = ExtractedDocument(source_id=source.source_id)

        try:
            from docx import Document

            document = Document(str(filepath))
            order = 0

            # Extract paragraphs
            for para in document.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Detect element type from style
                etype = ElementType.PARAGRAPH
                if para.style and para.style.name:
                    style_name = para.style.name.lower()
                    if "heading" in style_name or "title" in style_name:
                        etype = ElementType.HEADING
                    elif "list" in style_name:
                        etype = ElementType.LIST

                element = ExtractedElement(
                    source_id=source.source_id,
                    element_type=etype,
                    text=text,
                    order=order,
                )
                doc.elements.append(element)
                order += 1

            # Extract tables
            for table in document.tables:
                rows_data = []
                columns = []
                for i, row in enumerate(table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    if i == 0:
                        columns = cells
                    else:
                        rows_data.append(cells)

                ext_table = ExtractedTable(
                    source_id=source.source_id,
                    columns=columns if columns else [f"col_{j}" for j in range(len(table.columns))],
                    rows=rows_data,
                )
                doc.tables.append(ext_table)

                # Also add table as text element for pipeline
                table_elem = ExtractedElement(
                    source_id=source.source_id,
                    element_type=ElementType.TABLE,
                    text=ext_table.to_text(),
                    order=order,
                    table_id=ext_table.table_id,
                )
                doc.elements.append(table_elem)
                order += 1

            doc.raw_text = "\n\n".join(e.text for e in doc.elements)
            doc.diagnostics = ExtractionDiagnostics(
                total_characters=len(doc.raw_text),
                total_elements=len(doc.elements),
                tables_detected=len(doc.tables),
                status=ExtractionStatus.SUCCESS if doc.elements else ExtractionStatus.FAILED,
            )

        except Exception as e:
            logger.error("DOCX extraction failed", extra={"extra_data": {"source_id": source.source_id, "error": str(e)}})
            doc.diagnostics = ExtractionDiagnostics(status=ExtractionStatus.FAILED, issues=[f"DOCX error: {str(e)}"])

        return doc
